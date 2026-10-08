import os
import pandas as pd
import docx
import fitz

os.makedirs("sample_evidence", exist_ok=True)

# 1. Complete Evidence PDF
pdf_path = "sample_evidence/Complete_Access_Review_Q3_2026.pdf"
doc = fitz.open()
p = doc.new_page()
p.insert_text((50, 70), "PERIODIC USER ACCESS REVIEW — Q3 2026\nControl Code: C001\nReview Period: 2026-07-01 to 2026-09-30\nScope: IT Operations Team\n\n1. Access Review Report:\nListing of all 42 active privileged user accounts audited from Active Directory.\n\n2. Reviewer Confirmation:\nConfirmed and attested by Abhishek on 2026-09-28.\n\n3. Approval Evidence:\nFormal department head approval signed off via Change Ticket #CR-9402 on 2026-09-29.\n\n4. Exception Report:\nZero unauthorized privilege escalations. 2 dormant accounts revoked per SLA.")
doc.save(pdf_path)
doc.close()
print("Created Complete PDF")

# 2. Incomplete Evidence Excel (Lacks Exception Report and Formal Approval)
xlsx_path = "sample_evidence/Incomplete_User_Roster.xlsx"
df_users = pd.DataFrame({
    "User_ID": ["USR-01", "USR-02", "USR-03", "USR-04"],
    "Full_Name": ["John Doe", "Jane Smith", "Bob Vance", "Alice Cooper"],
    "Role": ["System Admin", "DBA", "Operator", "Security Analyst"],
    "Status": ["Active", "Active", "Active", "Active"]
})
df_users.to_excel(xlsx_path, sheet_name="Access_Review_Report", index=False)
print("Created Incomplete Excel")

# 3. CSV Evidence
csv_path = "sample_evidence/System_Accounts_Export.csv"
df_users.to_csv(csv_path, index=False)
print("Created CSV")

# 4. DOCX Evidence (Production Change Management)
docx_path = "sample_evidence/Production_Release_CAB_Signoff.docx"
doc2 = docx.Document()
doc2.add_heading("Production Change Management Authorization — Release v2.4", 0)
doc2.add_paragraph("Control Code: C002 | Application: Payments Application")
doc2.add_heading("Change Request Ticket with Peer Review", level=1)
doc2.add_paragraph("PR #1084 reviewed and approved by 2 Senior Software Engineers.")
doc2.add_heading("Automated Test Results Summary", level=1)
doc2.add_paragraph("Jenkins CI run #412 passed 100% of unit and integration tests (342/342).")
doc2.add_heading("Production Release Authorization Sign-off", level=1)
doc2.add_paragraph("Approved by CAB Lead and Director of Engineering for deployment.")
doc2.save(docx_path)
print("Created DOCX")
