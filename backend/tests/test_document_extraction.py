import io
import pandas as pd
from app.services.documents.extractor import document_extractor
from app.services.storage.storage_service import storage_service

def test_sha256_calculation():
    data = b"Sample LOD2 evidence bytes 2026"
    h1 = storage_service.calculate_sha256(data)
    h2 = storage_service.calculate_sha256(data)
    assert h1 == h2
    assert len(h1) == 64

def test_csv_extraction():
    csv_bytes = b"EmployeeID,Name,Role,Status\n101,Alice,Reviewer,Active\n102,Bob,Developer,Active\n"
    text = document_extractor.extract_text(csv_bytes, "access_log.csv", ".csv")
    assert "EmployeeID" in text
    assert "Alice" in text
    assert "Bob" in text

def test_excel_extraction():
    output = io.BytesIO()
    df = pd.DataFrame({"Username": ["admin", "guest"], "Privilege": ["Root", "ReadOnly"]})
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Accounts", index=False)
    excel_bytes = output.getvalue()

    text = document_extractor.extract_text(excel_bytes, "accounts.xlsx", ".xlsx")
    assert "=== SHEET: Accounts ===" in text
    assert "admin" in text
    assert "ReadOnly" in text

def test_docx_extraction():
    try:
        import docx
        doc = docx.Document()
        doc.add_heading("User Access Review Sign-Off", level=1)
        doc.add_paragraph("All privileges reviewed and verified by CISO.")
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        text = document_extractor.extract_text(docx_bytes, "signoff.docx", ".docx")
        assert "User Access Review Sign-Off" in text
        assert "All privileges reviewed" in text
    except ImportError:
        pass

def test_pdf_extraction():
    try:
        import fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 72), "LOD2 Audit Test Document - Access Approval Confirmed")
        pdf_bytes = doc.write()
        doc.close()

        text = document_extractor.extract_text(pdf_bytes, "approval.pdf", ".pdf")
        assert "LOD2 Audit Test Document" in text
        assert "Access Approval Confirmed" in text
    except ImportError:
        pass
