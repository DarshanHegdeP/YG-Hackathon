import io
import os
import pandas as pd
from typing import Optional

class DocumentExtractor:
    @staticmethod
    def extract_text(file_bytes: bytes, file_name: str, file_type: str) -> str:
        """
        Extracts structured text from PDF, XLSX, DOCX, and CSV files.
        """
        ext = os.path.splitext(file_name)[1].lower() if file_name else ""
        if not ext and file_type:
            ext = "." + file_type.lower().lstrip(".")

        try:
            if ext == ".pdf":
                return DocumentExtractor._extract_pdf(file_bytes)
            elif ext in [".xlsx", ".xls"]:
                return DocumentExtractor._extract_excel(file_bytes)
            elif ext == ".docx":
                return DocumentExtractor._extract_docx(file_bytes)
            elif ext == ".csv":
                return DocumentExtractor._extract_csv(file_bytes)
            elif ext in [".txt", ".json", ".log"]:
                return file_bytes.decode("utf-8", errors="replace")
            else:
                return f"[Unsupported file extension {ext} for deep parsing. Raw size: {len(file_bytes)} bytes]"
        except Exception as e:
            return f"[Extraction Error: Failed to extract text from {file_name}: {str(e)}]"

    @staticmethod
    def _extract_pdf(file_bytes: bytes) -> str:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        text_parts = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text()
            if text.strip():
                text_parts.append(f"--- PAGE {page_num + 1} ---\n{text.strip()}")
        doc.close()
        return "\n\n".join(text_parts) if text_parts else "[PDF contains no readable text]"

    @staticmethod
    def _extract_excel(file_bytes: bytes) -> str:
        excel_file = io.BytesIO(file_bytes)
        xls = pd.ExcelFile(excel_file)
        text_parts = []
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name)
            text_parts.append(f"=== SHEET: {sheet_name} ===")
            if df.empty:
                text_parts.append("(Sheet is empty)")
            else:
                # Include column headers and data formatted
                headers = list(df.columns)
                text_parts.append(f"Columns: {', '.join([str(h) for h in headers])}")
                text_parts.append(f"Row count: {len(df)}")
                # Format up to 100 rows cleanly
                sample_df = df.head(100)
                text_parts.append(sample_df.to_string(index=False))
        return "\n\n".join(text_parts)

    @staticmethod
    def _extract_docx(file_bytes: bytes) -> str:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        text_parts = []

        # Extract paragraphs
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        if paragraphs:
            text_parts.append("=== DOCUMENT PARAGRAPHS ===")
            text_parts.extend(paragraphs)

        # Extract tables
        if doc.tables:
            text_parts.append("\n=== DOCUMENT TABLES ===")
            for table_idx, table in enumerate(doc.tables):
                table_lines = [f"[Table {table_idx + 1}]"]
                for row in table.rows:
                    row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    table_lines.append(" | ".join(row_data))
                text_parts.append("\n".join(table_lines))

        return "\n\n".join(text_parts) if text_parts else "[DOCX contains no text or tables]"

    @staticmethod
    def _extract_csv(file_bytes: bytes) -> str:
        try:
            df = pd.read_csv(io.BytesIO(file_bytes))
        except UnicodeDecodeError:
            df = pd.read_csv(io.BytesIO(file_bytes), encoding="latin1")

        text_parts = ["=== CSV DATA ==="]
        text_parts.append(f"Columns: {', '.join([str(c) for c in df.columns])}")
        text_parts.append(f"Total Rows: {len(df)}")
        text_parts.append(df.head(100).to_string(index=False))
        return "\n\n".join(text_parts)

document_extractor = DocumentExtractor()
